from pptx import Presentation
from PIL import Image, ImageDraw
from src.ports.image_ports.pptx_to_image_port import PptxToImagePort
import logging

# Configurar el logger
logger = logging.getLogger(__name__)

class PptxToImageAdapter(PptxToImagePort):
    def __init__(self, file_path):
        self.file_path = file_path
        logger.info("PptToImageAdapter inicializado con la ruta del archivo")

    def convert_to_image(self):
        logger.debug("Iniciando conversión de presentación a imagen")
        try:
            # Abrir la presentación desde la ruta del archivo
            presentation = Presentation(self.file_path)
            logger.info(f"Procesando {len(presentation.slides)} diapositivas")
            images = []
            
            for i, slide in enumerate(presentation.slides, 1):
                logger.debug(f"Procesando diapositiva {i}")
                img = Image.new('RGB', (1280, 720), color=(255, 255, 255))
                d = ImageDraw.Draw(img)
                slide_text = "\n".join([shape.text for shape in slide.shapes if hasattr(shape, "text")])
                d.text((10, 10), slide_text, fill=(0, 0, 0))
                images.append(img)
                logger.debug(f"Diapositiva {i} procesada exitosamente")

            logger.info("Conversión completada exitosamente")
            return images
            
        except Exception as e:
            logger.error(f"Error durante la conversión: {str(e)}")
            raise
